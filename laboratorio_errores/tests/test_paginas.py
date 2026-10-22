"""Pruebas de humo de la interfaz con streamlit.testing (sin navegador).

Cada página se ejecuta, se mueven los controles principales y se comprueba que no haya
excepciones ni mensajes de error.
"""

import json
from pathlib import Path

import pytest

pytest.importorskip("streamlit")
from streamlit.testing.v1 import AppTest  # noqa: E402

RAIZ = Path(__file__).resolve().parents[1]
T = 120


def _pagina(nombre: str) -> AppTest:
    at = AppTest.from_file(str(RAIZ / "paginas" / nombre), default_timeout=T)
    at.run()
    _sano(at)
    return at


def _sano(at: AppTest) -> None:
    assert not at.exception, at.exception
    assert not at.error, [e.value for e in at.error]


def _boton(at: AppTest, inicio: str) -> AppTest:
    for b in at.button:
        if b.label.startswith(inicio):
            b.click()
            at.run()
            _sano(at)
            return at
    raise AssertionError(f"No encontré el botón «{inicio}»")


def test_app_completa_arranca():
    at = AppTest.from_file(str(RAIZ / "app.py"), default_timeout=T)
    at.run()
    _sano(at)


def test_inicio():
    at = _pagina("inicio.py")
    assert "0.00369" in " ".join(m.value for m in at.markdown)


def test_modulo1_caso_libro_y_controles():
    at = _pagina("m1_binomial.py")
    assert any("Coincide con el libro" in s.value for s in at.success)
    _boton(at, "Región que compara el libro")
    assert any("β(.3) ≈ .278" in s.value for s in at.success)
    at.slider(key="b_n").set_value(40).run()
    _sano(at)
    at.radio(key="b_cola").set_value("superior").run()
    _sano(at)
    assert at.session_state.b_p1 > at.session_state.b_p0
    at.slider(key="b_c").set_value(30).run()
    _sano(at)


def test_modulo1_p1_fuera_de_la_alternativa_avisa():
    at = _pagina("m1_binomial.py")
    at.slider(key="b_p1").set_value(0.7).run()
    assert not at.exception and at.error


def test_modulo2_frontera_manual_y_json():
    at = _pagina("m2_dilema.py")
    at.radio(key="z_modo").set_value("Muevo la frontera k a mano").run()
    _sano(at)
    at.slider(key="z_k").set_value(-0.1).run()
    _sano(at)
    at.radio(key="z_cola").set_value("superior").run()
    _sano(at)
    assert at.session_state.z_mu1 > at.session_state.z_mu0
    at.number_input(key="z_n").set_value(200).run()
    _sano(at)


def test_modulo2_json_exportado_es_valido():
    from nucleo import pruebaz as pz
    d = json.loads(pz.a_json(pz.ConfigZ()))
    assert {"mu0", "mu1", "sigma", "n", "alfa", "tipo_prueba"} <= set(d)


@pytest.mark.parametrize("prueba", ["Prueba Z (Módulo 2)", "Binomial de Jones (Módulo 1)"])
@pytest.mark.parametrize("escenario", ["Solo H0 verdadera (estimar α)", "Solo Ha verdadera (estimar β)",
                                       "Mitad H0, mitad Ha"])
def test_modulo3_simula(prueba, escenario):
    at = _pagina("m3_montecarlo.py")
    at.radio(key="mc_prueba").set_value(prueba)
    at.radio(key="mc_escenario").set_value(escenario)
    at.select_slider(key="mc_m").set_value(1000).run()
    _boton(at, "Simular")
    assert "mc_res" in at.session_state
    assert len(at.dataframe) >= 2


@pytest.mark.parametrize("prueba", ["Prueba Z (Módulo 2)", "Binomial de Jones (Módulo 1)"])
def test_modulo4_rondas(prueba):
    at = _pagina("m4_decision.py")
    at.radio(key="d_prueba").set_value(prueba).run()
    for decision in ("Rechazar H₀", "No rechazar H₀", "Rechazar H₀"):
        _boton(at, "Nuevo experimento")
        _boton(at, decision)
        texto = " ".join(m.value for m in at.markdown)
        assert "Estado verdadero usado para generar la muestra" in texto
    m = at.session_state.d_marcador
    assert m["rondas"] == 3 and m["correctas"] + m["tipo I"] + m["tipo II"] == 3
    _boton(at, "Reiniciar marcador")
    assert at.session_state.d_marcador["rondas"] == 0


def test_modulo5_signos():
    at = _pagina("m5_signos.py")
    at.select_slider(key="s_m").set_value(2000).run()
    _sano(at)
    at.radio(key="s_cola").set_value("inferior").run()
    _sano(at)


def test_excel_preparar_libro_con_monte_carlo():
    at = _pagina("excel.py")
    _boton(at, "Preparar el libro de Excel")
    assert "x_libro" in at.session_state
    from nucleo import montecarlo as mc
    from nucleo import pruebaz as pz
    at.session_state["mc_res"] = {"df": mc.simular_z(pz.ConfigZ(), 200, 0.5, 1), "params": pz.como_dict(pz.ConfigZ()),
                                  "prueba": "Prueba Z (Módulo 2)", "alfa": 0.05, "beta": 0.3, "escenario": "x"}
    at.run()
    _boton(at, "Preparar el libro de Excel")
    from io import BytesIO

    from openpyxl import load_workbook
    assert "Monte_Carlo" in load_workbook(BytesIO(at.session_state.x_libro)).sheetnames


def test_excel_analiza_datos_subidos():
    """AppTest no simula subidas: se carga directamente lo que dejaría el callback de subida."""
    from nucleo import excel as xl
    from nucleo import pruebaz as pz
    at = AppTest.from_file(str(RAIZ / "paginas" / "excel.py"), default_timeout=T)
    res = xl.leer(xl.plantilla(pz.ConfigZ()))
    at.session_state["x_res"] = res
    for k, v in res["params_datos"].items():
        at.session_state[f"x_{k}"] = v
    at.run()
    _sano(at)
    assert any("valor p" in m.value or "Decisión" in m.value for m in at.markdown)
    at.number_input(key="x_mu0").set_value(5.0).run()
    _sano(at)


def test_excel_columnas_de_un_csv():
    from nucleo import excel as xl
    at = AppTest.from_file(str(RAIZ / "paginas" / "excel.py"), default_timeout=T)
    res = xl.leer_csv("fecha;TRM;inflacion\n1;4000,5;0,5\n2;4010,1;0,7\n3;3990,2;0,6\n".encode())
    at.session_state["x_res"] = res
    at.session_state["x_col"] = "TRM"
    for k, v in {"mu0": 4000.0, "sigma": 10.0, "alfa": 0.05, "cola": "superior"}.items():
        at.session_state[f"x_{k}"] = v
    at.run()
    _sano(at)
    at.selectbox(key="x_col").set_value("inflacion").run()
    _sano(at)


def test_excel_usar_configuracion_subida():
    from nucleo import excel as xl
    from nucleo import pruebaz as pz
    cfg = pz.ConfigZ(10, 12, 3, 50, 0.01, "superior")
    libro = xl.construir_libro(dict(n=20, p0=0.5, p1=0.25, c=5, cola="inferior"), cfg)
    at = AppTest.from_file(str(RAIZ / "paginas" / "excel.py"), default_timeout=T)
    at.session_state["x_cfg"] = xl.leer(libro)
    at.run()
    _sano(at)
    _boton(at, "Usar estos valores en la app")
    s = at.session_state
    assert (s.z_mu0, s.z_mu1, s.z_n, s.z_cola) == (10, 12, 50, "superior")
    assert (s.b_n, s.b_c, s.b_p1) == (20, 5, 0.25)
    assert any("Listo" in m.value for m in at.success)
