"""Libro de Excel: estructura, fórmulas y lectura de archivos subidos."""

import io

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from nucleo import excel as xl
from nucleo import montecarlo as mc
from nucleo import pruebaz as pz

JONES = dict(n=15, p0=0.5, p1=0.3, c=2, cola="inferior")


def _libro(**kw):
    cfg = kw.pop("cfg", pz.ConfigZ())
    return xl.construir_libro(JONES, cfg, **kw), cfg


def test_hojas_y_formulas():
    datos, cfg = _libro(mc=mc.simular_z(pz.ConfigZ(), 300, 0.5, 1))
    wb = load_workbook(io.BytesIO(datos))
    assert wb.sheetnames == ["Leeme", "Binomial", "Curva_beta", "Prueba_Z", "Monte_Carlo", "Tus_datos"]
    b = wb["Binomial"]
    assert b["B11"].value.startswith("=") and "BINOMDIST" in b["B11"].value
    assert b["B12"].value.startswith("=1-")
    z = wb["Prueba_Z"]
    assert "NORMSINV" in z["B13"].value and "ROUNDUP" in z["B22"].value
    m = wb["Monte_Carlo"]
    assert "COUNTIFS" in m["F5"].value
    assert m["A22"].value == 1 and m["D22"].value.startswith("=")


def test_sin_monte_carlo_no_hay_hoja():
    datos, _ = _libro()
    assert "Monte_Carlo" not in load_workbook(io.BytesIO(datos)).sheetnames


def test_solo_funciones_compatibles():
    """Nada de funciones que pidan prefijo _xlfn o que LibreOffice no evalúe."""
    datos, _ = _libro(mc=mc.simular_z(pz.ConfigZ(), 50, 0.5, 1))
    wb = load_workbook(io.BytesIO(datos))
    prohibidas = ("BINOM.DIST", "NORM.S.INV", "NORM.DIST", "XLOOKUP", "FILTER", "SEQUENCE", "_xlfn")
    for ws in wb.worksheets:
        for fila in ws.iter_rows():
            for c in fila:
                if isinstance(c.value, str) and c.value.startswith("="):
                    assert not any(p in c.value for p in prohibidas), (ws.title, c.coordinate, c.value)


def test_ida_y_vuelta_de_parametros():
    cfg = pz.ConfigZ(10, 11.5, 4, 40, 0.01, "superior")
    datos, _ = _libro(cfg=cfg)
    r = xl.leer(datos)
    assert r["config_z"] == cfg
    assert r["binomial"] == JONES
    assert len(r["datos"]) == 40
    assert r["params_datos"] == {"mu0": 10, "sigma": 4, "alfa": 0.01, "cola": "superior"}


def test_parametros_editados_en_excel():
    datos, _ = _libro()
    wb = load_workbook(io.BytesIO(datos))
    wb["Binomial"]["B7"] = 5
    wb["Prueba_Z"]["B7"] = 100
    wb["Prueba_Z"]["B9"] = "Superior"          # mayúscula: también vale
    wb["Prueba_Z"]["B5"] = 0.3
    buf = io.BytesIO()
    wb.save(buf)
    r = xl.leer(buf.getvalue())
    assert r["binomial"]["c"] == 5
    assert r["config_z"] == pz.ConfigZ(0.0, 0.3, 1.0, 100, 0.05, "superior")


def test_parametros_invalidos_dan_aviso():
    datos, _ = _libro()
    wb = load_workbook(io.BytesIO(datos))
    wb["Prueba_Z"]["B5"] = 5          # con Hₐ: μ < μ0, μ1 = 5 no está en la alternativa
    wb["Binomial"]["B6"] = 0.8
    buf = io.BytesIO()
    wb.save(buf)
    r = xl.leer(buf.getvalue())
    assert "config_z" not in r and "binomial" not in r
    assert len(r["avisos"]) == 2


def test_excel_cualquiera_con_varias_columnas():
    wb = Workbook()
    ws = wb.active
    ws.append(["fecha", "TRM", "inflacion"])
    for i in range(10):
        ws.append([f"2024-01-{i + 1:02d}", 4000 + i, 0.5 + i / 10])
    buf = io.BytesIO()
    wb.save(buf)
    r = xl.leer(buf.getvalue())
    assert r["columnas"] == ["TRM", "inflacion"]
    assert np.allclose(r["datos"], 4000 + np.arange(10))
    assert np.allclose(xl.columna(r["tabla"], "inflacion"), 0.5 + np.arange(10) / 10)


def test_csv_con_punto_y_coma_y_coma_decimal():
    texto = "fecha;valor\n2024-01;4.012,50\n2024-02;3.990,25\n2024-03;4.101,00\n"
    r = xl.leer_csv(texto.encode("latin-1"))
    assert r["columnas"] == ["valor"]
    assert np.allclose(xl.columna(r["tabla"], "valor"), [4012.5, 3990.25, 4101.0])


def test_csv_normal():
    df = pd.DataFrame({"x": np.arange(5.0), "nombre": list("abcde")})
    r = xl.leer_csv(df.to_csv(index=False).encode())
    assert r["columnas"] == ["x"] and len(r["datos"]) == 5


def test_archivo_que_no_es_excel():
    with pytest.raises(ValueError):
        xl.leer(b"esto no es un xlsx")


def test_plantilla():
    r = xl.leer(xl.plantilla(pz.ConfigZ()))
    assert "datos" in r and "config_z" not in r
