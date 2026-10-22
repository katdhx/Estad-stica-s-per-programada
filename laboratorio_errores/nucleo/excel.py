"""Puente con Excel: libro con fórmulas vivas y lectura de archivos subidos.

El libro que se descarga NO trae números pegados: trae fórmulas de Excel
(BINOMDIST, NORMSINV, NORMDIST, COUNTIFS…) que recalculan α, β y la potencia cuando se
cambia una celda amarilla. Así el profesor puede comprobar en Excel, de forma
independiente, lo que muestra la aplicación.

Se usan las funciones «de compatibilidad» (BINOMDIST, NORMSINV, NORMDIST, NORMSDIST)
porque abren sin errores en cualquier versión de Excel y en LibreOffice. Excel en español
las muestra como DISTR.BINOM, DISTR.NORM.ESTAND.INV, DISTR.NORM y DISTR.NORM.ESTAND.

Hojas del libro
---------------
Leeme        instrucciones.
Binomial     el ejemplo de Jones: entradas n, p0, p1, c, cola; tabla de probabilidades,
             región de rechazo, α, β y potencia.
Curva_beta   β(p) y potencia para p entre .01 y .99, con gráfico.
Prueba_Z     μ0, μ1, σ, n, α, cola → EE, z_α, frontera k, α, β, potencia y n requerido.
Monte_Carlo  (si hay simulación) las medias simuladas por la app; Excel vuelve a aplicar la
             regla de rechazo y cuenta los errores con COUNTIFS.
Tus_datos    plantilla: pega tus observaciones y Excel hace la prueba Z.
"""

from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from . import binomial as bn
from . import pruebaz as pz

N_MAX_TABLA = 100          # filas de la tabla binomial (y = 0, …, 100)
MC_MAX_FILAS = 10000       # experimentos Monte Carlo que se copian al libro
DATOS_MAX = 5000           # filas para observaciones en la plantilla

FUENTE = "Arial"
AZUL = "1B2A49"
AMARILLO = PatternFill("solid", fgColor="FFF2B3")
CABECERA = PatternFill("solid", fgColor=AZUL)
GRIS = PatternFill("solid", fgColor="EEF1F5")
BORDE = Border(bottom=Side(style="thin", color="C9D0DA"))


def _f(bold=False, color="1F2933", size=10, italic=False) -> Font:
    return Font(name=FUENTE, bold=bold, color=color, size=size, italic=italic)


def _titulo(ws, texto: str, sub: str) -> None:
    ws["A1"] = texto
    ws["A1"].font = _f(True, AZUL, 14)
    ws["A2"] = sub
    ws["A2"].font = _f(italic=True, color="55606E")


def _cabecera(ws, fila: int, textos: list[str], col0: int = 1) -> None:
    for j, t in enumerate(textos):
        c = ws.cell(row=fila, column=col0 + j, value=t)
        c.font = _f(True, "FFFFFF")
        c.fill = CABECERA
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _entrada(ws, celda: str, valor, fmt: str | None = None) -> None:
    ws[celda] = valor
    ws[celda].fill = AMARILLO
    ws[celda].font = _f(color="0000FF", bold=True)
    if fmt:
        ws[celda].number_format = fmt


def _formula(ws, celda: str, f: str, fmt: str = "0.0000", bold: bool = False) -> None:
    ws[celda] = f
    ws[celda].font = _f(bold=bold)
    ws[celda].number_format = fmt


def _etiqueta(ws, celda: str, texto: str, bold: bool = False) -> None:
    ws[celda] = texto
    ws[celda].font = _f(bold=bold)


def _lista(ws, celda: str, opciones: tuple[str, ...]) -> None:
    dv = DataValidation(type="list", formula1='"' + ",".join(opciones) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(celda)


# ------------------------------------------------------------------------------------ hojas

def _hoja_leeme(wb: Workbook, con_mc: bool) -> None:
    ws = wb.active
    ws.title = "Leeme"
    _titulo(ws, "El costo de equivocarse: errores tipo I y tipo II",
            "Libro de comprobación generado por la aplicación (Wackerly, Mendenhall y Scheaffer, 7.ª ed., cap. 10)")
    lineas = [
        ("Cómo se usa", True),
        ("1. Las celdas AMARILLAS con número azul son entradas: cámbialas y Excel recalcula todo.", False),
        ("2. Las celdas blancas son fórmulas de Excel, no números copiados de Python.", False),
        ("3. Compara los resultados con la aplicación: deben coincidir (salvo redondeo).", False),
        ("4. Puedes volver a subir este libro a la aplicación (página «Excel»): la app lee las", False),
        ("   entradas de las hojas Binomial y Prueba_Z, y las observaciones de la hoja Tus_datos.", False),
        ("", False),
        ("Hojas", True),
        ("Binomial: ejemplo de Jones. α = P(Y ∈ RR | p0), β = P(Y ∉ RR | p1), potencia = 1 − β.", False),
        ("Curva_beta: β(p) para muchos valores verdaderos p (toda la alternativa compuesta).", False),
        ("Prueba_Z: prueba Z para una media con σ conocida; frontera k, α, β, potencia y n requerido.", False),
    ]
    if con_mc:
        lineas.append(("Monte_Carlo: medias simuladas por la app; Excel aplica la regla y cuenta errores.", False))
    lineas += [
        ("Tus_datos: pega una columna de observaciones y Excel hace la prueba Z.", False),
        ("", False),
        ("Funciones usadas (nombre en Excel en español)", True),
        ("BINOMDIST = DISTR.BINOM · NORMSINV = DISTR.NORM.ESTAND.INV · NORMDIST = DISTR.NORM ·", False),
        ("NORMSDIST = DISTR.NORM.ESTAND · COUNTIFS = CONTAR.SI.CONJUNTO", False),
    ]
    for i, (t, b) in enumerate(lineas, start=4):
        ws.cell(row=i, column=1, value=t).font = _f(bold=b, color=AZUL if b else "1F2933")
    ws.column_dimensions["A"].width = 100


def _hoja_binomial(wb: Workbook, n: int, p0: float, p1: float, c: int, cola: str) -> None:
    ws = wb.create_sheet("Binomial")
    _titulo(ws, "Módulo 1 · El ejemplo binomial (Jones)",
            "Y ~ Binomial(n, p). Cola inferior: Ha: p < p0, RR = {y ≤ c}. Cola superior: Ha: p > p0, RR = {y ≥ c}.")
    _etiqueta(ws, "A3", "Entradas", True)
    filas = [("n (votantes)", n, "0"), ("p0 (bajo H0)", p0, "0.000"), ("p1 (alternativa específica)", p1, "0.000"),
             ("c (umbral de rechazo)", c, "0"), ("cola", cola, None)]
    for i, (et, v, fmt) in enumerate(filas, start=4):
        _etiqueta(ws, f"A{i}", et)
        _entrada(ws, f"B{i}", v, fmt)
    _lista(ws, "B8", bn.COLAS)
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2=str(N_MAX_TABLA))
    ws.add_data_validation(dv)
    dv.add("B4")

    _etiqueta(ws, "A10", "Resultados", True)
    rech = 'IF($B$8="inferior",BINOMDIST($B$7,$B$4,{p},TRUE),IF($B$7<=0,1,1-BINOMDIST($B$7-1,$B$4,{p},TRUE)))'
    _etiqueta(ws, "A11", "α = P(Y ∈ RR | p0)")
    _formula(ws, "B11", "=" + rech.format(p="$B$5"), bold=True)
    _etiqueta(ws, "A12", "β(p1) = P(Y ∉ RR | p1)")
    _formula(ws, "B12", "=1-" + rech.format(p="$B$6"), bold=True)
    _etiqueta(ws, "A13", "Potencia en p1 = 1 − β(p1)")
    _formula(ws, "B13", "=1-B12", bold=True)

    _cabecera(ws, 15, ["y", "P(Y = y | p0)", "P(Y = y | p1)", "Región"])
    for k in range(N_MAX_TABLA + 1):
        r = 16 + k
        ws.cell(row=r, column=1, value=k).font = _f()
        ws.cell(row=r, column=2, value=f'=IF(A{r}>$B$4,"",BINOMDIST(A{r},$B$4,$B$5,FALSE))').number_format = "0.00000"
        ws.cell(row=r, column=3, value=f'=IF(A{r}>$B$4,"",BINOMDIST(A{r},$B$4,$B$6,FALSE))').number_format = "0.00000"
        ws.cell(row=r, column=4, value=(f'=IF(A{r}>$B$4,"",IF(OR(AND($B$8="inferior",A{r}<=$B$7),'
                                         f'AND($B$8="superior",A{r}>=$B$7)),"Rechazo","No rechazo"))'))
        for col in (2, 3, 4):
            ws.cell(row=r, column=col).font = _f()
    ultima = 16 + n
    _etiqueta(ws, "F15", "Suma de control (debe dar 1):")
    _formula(ws, "F16", f"=SUM(B16:B{16 + N_MAX_TABLA})", "0.000000")

    graf = BarChart()
    graf.type = "col"
    graf.title = "Distribuciones de Y bajo p0 y bajo p1"
    graf.y_axis.title = "Probabilidad"
    graf.x_axis.title = "y"
    datos = Reference(ws, min_col=2, max_col=3, min_row=15, max_row=ultima)
    graf.add_data(datos, titles_from_data=True)
    graf.set_categories(Reference(ws, min_col=1, min_row=16, max_row=ultima))
    graf.height, graf.width = 8, 16
    ws.add_chart(graf, "F18")
    for col, w in zip("ABCDEF", (30, 16, 16, 13, 3, 30)):
        ws.column_dimensions[col].width = w


def _hoja_curva(wb: Workbook) -> None:
    ws = wb.create_sheet("Curva_beta")
    _titulo(ws, "β(p) para toda la alternativa compuesta",
            "Usa n, c y la cola de la hoja Binomial. Solo los p que están en Ha son errores tipo II.")
    _cabecera(ws, 4, ["p verdadero", "β(p)", "Potencia 1 − β(p)", "¿p en Ha?"])
    rech = ('IF(Binomial!$B$8="inferior",BINOMDIST(Binomial!$B$7,Binomial!$B$4,A{r},TRUE),'
            'IF(Binomial!$B$7<=0,1,1-BINOMDIST(Binomial!$B$7-1,Binomial!$B$4,A{r},TRUE)))')
    for i in range(99):
        r = 5 + i
        ws.cell(row=r, column=1, value=round((i + 1) / 100, 2)).number_format = "0.00"
        ws.cell(row=r, column=2, value="=1-" + rech.format(r=r)).number_format = "0.0000"
        ws.cell(row=r, column=3, value=f"=1-B{r}").number_format = "0.0000"
        ws.cell(row=r, column=4, value=(f'=IF(OR(AND(Binomial!$B$8="inferior",A{r}<Binomial!$B$5),'
                                         f'AND(Binomial!$B$8="superior",A{r}>Binomial!$B$5)),"sí","no")'))
        for col in range(1, 5):
            ws.cell(row=r, column=col).font = _f()
    graf = ScatterChart()
    graf.title = "Curva β(p) y potencia"
    graf.x_axis.title = "p verdadero"
    graf.y_axis.title = "Probabilidad"
    graf.style = 13
    x = Reference(ws, min_col=1, min_row=5, max_row=103)
    for col in (2, 3):
        s = Series(Reference(ws, min_col=col, min_row=4, max_row=103), x, title_from_data=True)
        s.marker.symbol = "none"
        s.smooth = True
        graf.series.append(s)
    graf.height, graf.width = 9, 16
    ws.add_chart(graf, "F4")
    for col, w in zip("ABCD", (14, 12, 18, 11)):
        ws.column_dimensions[col].width = w


def _hoja_z(wb: Workbook, cfg: pz.ConfigZ, beta_obj: float) -> None:
    ws = wb.create_sheet("Prueba_Z")
    _titulo(ws, "Módulo 2 · Prueba Z para una media (σ conocida)",
            "Superior: Ha: μ > μ0, RR = {ȳ > k}. Inferior: Ha: μ < μ0, RR = {ȳ < k}. k = μ0 ± z_α·σ/√n.")
    _etiqueta(ws, "A3", "Entradas", True)
    filas = [("μ0 (media bajo H0)", cfg.mu0, "0.000"), ("μ1 (alternativa específica)", cfg.mu1, "0.000"),
             ("σ (conocida)", cfg.sigma, "0.000"), ("n", int(cfg.n), "0"), ("α", cfg.alfa, "0.000"),
             ("cola", cfg.cola, None)]
    for i, (et, v, fmt) in enumerate(filas, start=4):
        _etiqueta(ws, f"A{i}", et)
        _entrada(ws, f"B{i}", v, fmt)
    _lista(ws, "B9", pz.COLAS)

    _etiqueta(ws, "A11", "Resultados", True)
    res = [
        ("EE = σ/√n", "=B6/SQRT(B7)"),
        ("z_α  (P(Z > z_α) = α)", "=NORMSINV(1-B8)"),
        ("Frontera k", '=IF(B9="superior",B4+B13*B12,B4-B13*B12)'),
        ("α comprobado = P(Ȳ ∈ RR | μ0)", '=IF(B9="superior",1-NORMDIST(B14,B4,B12,TRUE),NORMDIST(B14,B4,B12,TRUE))'),
        ("β(μ1) = P(Ȳ ∉ RR | μ1)", '=IF(B9="superior",NORMDIST(B14,B5,B12,TRUE),1-NORMDIST(B14,B5,B12,TRUE))'),
        ("Potencia en μ1 = 1 − β", "=1-B16"),
    ]
    for i, (et, f) in enumerate(res, start=12):
        _etiqueta(ws, f"A{i}", et)
        _formula(ws, f"B{i}", f, bold=i >= 15)

    _etiqueta(ws, "A19", "Tamaño muestral (Sección 10.4)", True)
    _etiqueta(ws, "A20", "β deseado en μ1")
    _entrada(ws, "B20", beta_obj, "0.000")
    _etiqueta(ws, "A21", "z_β")
    _formula(ws, "B21", "=NORMSINV(1-B20)")
    _etiqueta(ws, "A22", "n = (z_α + z_β)² σ² / (μ1 − μ0)²")
    _formula(ws, "B22", "=ROUNDUP((B13+B21)^2*B6^2/(B5-B4)^2,0)", "0", bold=True)

    _etiqueta(ws, "D3", "Mover la frontera a mano", True)
    _etiqueta(ws, "D4", "k propuesto")
    _entrada(ws, "E4", round(pz.frontera(cfg), 4), "0.0000")
    _etiqueta(ws, "D5", "α(k)")
    _formula(ws, "E5", '=IF(B9="superior",1-NORMDIST(E4,B4,B12,TRUE),NORMDIST(E4,B4,B12,TRUE))')
    _etiqueta(ws, "D6", "β(k)")
    _formula(ws, "E6", '=IF(B9="superior",NORMDIST(E4,B5,B12,TRUE),1-NORMDIST(E4,B5,B12,TRUE))')
    ws["D8"] = "Si k se aleja de μ0 hacia Ha, α(k) baja y β(k) sube."
    ws["D8"].font = _f(italic=True, color="55606E")
    for col, w in zip("ABCDE", (36, 14, 3, 22, 12)):
        ws.column_dimensions[col].width = w


def _hoja_mc(wb: Workbook, df: pd.DataFrame, cfg: pz.ConfigZ) -> None:
    ws = wb.create_sheet("Monte_Carlo")
    _titulo(ws, "Módulo 3 · Monte Carlo (prueba Z)",
            "Las medias vienen de la app. La decisión y el tipo de error los calcula Excel con las entradas de abajo.")
    df = df.head(MC_MAX_FILAS)
    m = len(df)
    _etiqueta(ws, "A3", "Parámetros usados", True)
    for i, (et, v, fmt) in enumerate([("μ0", cfg.mu0, "0.000"), ("μ1", cfg.mu1, "0.000"), ("σ", cfg.sigma, "0.000"),
                                      ("n", int(cfg.n), "0"), ("α", cfg.alfa, "0.000"), ("cola", cfg.cola, None)],
                                     start=4):
        _etiqueta(ws, f"A{i}", et)
        _entrada(ws, f"B{i}", v, fmt)
    _lista(ws, "B9", pz.COLAS)
    _etiqueta(ws, "A10", "z_α")
    _formula(ws, "B10", "=NORMSINV(1-B8)")

    ini, fin = 22, 22 + m - 1
    rng_estado, rng_res = f"B{ini}:B{fin}", f"F{ini}:F{fin}"
    _etiqueta(ws, "D3", "Resumen calculado por Excel", True)
    _cabecera(ws, 4, ["", "Experimentos", "Errores", "Frecuencia", "Teórica"], col0=4)
    _etiqueta(ws, "D5", "H0 verdadera → error tipo I")
    ws["E5"] = f'=COUNTIF({rng_estado},"H0 verdadera")'
    ws["F5"] = f'=COUNTIFS({rng_estado},"H0 verdadera",{rng_res},"Error tipo I")'
    ws["G5"] = "=IF(E5=0,\"\",F5/E5)"
    ws["H5"] = "=B8"
    _etiqueta(ws, "D6", "Ha verdadera → error tipo II")
    ws["E6"] = f'=COUNTIF({rng_estado},"Ha verdadera")'
    ws["F6"] = f'=COUNTIFS({rng_estado},"Ha verdadera",{rng_res},"Error tipo II")'
    ws["G6"] = "=IF(E6=0,\"\",F6/E6)"
    ws["H6"] = ('=IF(B9="superior",NORMDIST(B4+B10*B6/SQRT(B7),B5,B6/SQRT(B7),TRUE),'
                '1-NORMDIST(B4-B10*B6/SQRT(B7),B5,B6/SQRT(B7),TRUE))')
    _etiqueta(ws, "D7", "Error Monte Carlo de la frecuencia (tipo I)")
    ws["G7"] = '=IF(E5=0,"",SQRT(G5*(1-G5)/E5))'
    _etiqueta(ws, "D8", "Error Monte Carlo de la frecuencia (tipo II)")
    ws["G8"] = '=IF(E6=0,"",SQRT(G6*(1-G6)/E6))'
    for c in ("G5", "G6", "H5", "H6", "G7", "G8"):
        ws[c].number_format = "0.0000"
    for fila in range(5, 9):
        for col in "DEFGH":
            ws[f"{col}{fila}"].font = _f(bold=col == "G")
    if m < 1:
        return
    _cabecera(ws, 21, ["Experimento", "Estado verdadero", "Media muestral ȳ", "z", "Decisión", "Resultado"])
    for i, fila in enumerate(df.itertuples(index=False)):
        r = ini + i
        ws.cell(row=r, column=1, value=int(fila.experimento))
        ws.cell(row=r, column=2, value=fila.estado)
        ws.cell(row=r, column=3, value=float(fila.media_muestral)).number_format = "0.0000"
        ws.cell(row=r, column=4, value=f"=(C{r}-$B$4)/($B$6/SQRT($B$7))").number_format = "0.0000"
        ws.cell(row=r, column=5, value=(f'=IF(OR(AND($B$9="superior",D{r}>$B$10),AND($B$9="inferior",D{r}<-$B$10)),'
                                         f'"Rechazar H0","No rechazar H0")'))
        ws.cell(row=r, column=6, value=(f'=IF(AND(B{r}="H0 verdadera",E{r}="Rechazar H0"),"Error tipo I",'
                                         f'IF(AND(B{r}="Ha verdadera",E{r}="No rechazar H0"),"Error tipo II",'
                                         f'"Decisión correcta"))'))
    ws.freeze_panes = "A22"
    for col, w in zip("ABCDEFGH", (13, 34, 16, 12, 16, 18, 12, 10)):
        ws.column_dimensions[col].width = w


def _hoja_datos(wb: Workbook, cfg: pz.ConfigZ, ejemplo: np.ndarray | None) -> None:
    ws = wb.create_sheet("Tus_datos")
    ws["A1"] = "observacion"
    ws["A1"].font = _f(True, "FFFFFF")
    ws["A1"].fill = CABECERA
    if ejemplo is not None:
        for i, v in enumerate(ejemplo, start=2):
            c = ws.cell(row=i, column=1, value=round(float(v), 3))
            c.fill = AMARILLO
            c.font = _f(color="0000FF")
    rango = f"A2:A{DATOS_MAX + 1}"
    _etiqueta(ws, "C1", "Pega tus observaciones en la columna A (una por fila, desde A2).", True)
    _etiqueta(ws, "C2", "Los valores de ejemplo son simulados: bórralos y pon los tuyos.")
    ws["C2"].font = _f(italic=True, color="55606E")
    for i, (et, v, fmt) in enumerate([("μ0 (media bajo H0)", cfg.mu0, "0.000"), ("σ (conocida)", cfg.sigma, "0.000"),
                                      ("α", cfg.alfa, "0.000"), ("cola", cfg.cola, None)], start=4):
        _etiqueta(ws, f"C{i}", et)
        _entrada(ws, f"D{i}", v, fmt)
    _lista(ws, "D7", pz.COLAS)
    _etiqueta(ws, "C9", "Resultados (fórmulas de Excel)", True)
    res = [
        ("n", f"=COUNT({rango})", "0"),
        ("ȳ", f'=IF(D10=0,"",AVERAGE({rango}))', "0.0000"),
        ("EE = σ/√n", '=IF(D10=0,"",D5/SQRT(D10))', "0.0000"),
        ("z = (ȳ − μ0)/EE", '=IF(D10=0,"",(D11-D4)/D12)', "0.0000"),
        ("z_α", "=NORMSINV(1-D6)", "0.0000"),
        ("Decisión", ('=IF(D10=0,"",IF(OR(AND(D7="superior",D13>D14),AND(D7="inferior",D13<-D14)),'
                      '"Rechazar H0","No rechazar H0"))'), None),
        ("Valor p", '=IF(D10=0,"",IF(D7="superior",1-NORMSDIST(D13),NORMSDIST(D13)))', "0.0000"),
    ]
    for i, (et, f, fmt) in enumerate(res, start=10):
        _etiqueta(ws, f"C{i}", et)
        _formula(ws, f"D{i}", f, fmt or "General", bold=et in ("Decisión", "Valor p"))
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 3
    ws.column_dimensions["C"].width = 30
    ws.column_dimensions["D"].width = 18


def construir_libro(binom: dict, cfg: pz.ConfigZ, beta_obj: float = 0.10, mc: pd.DataFrame | None = None,
                    ejemplo_datos: np.ndarray | None = None) -> bytes:
    """Devuelve el .xlsx en bytes. ``binom`` = dict(n, p0, p1, c, cola)."""
    bn.validar(binom["n"], binom["p0"], binom["c"], binom["cola"])
    if binom["n"] > N_MAX_TABLA:
        raise ValueError(f"La hoja Binomial admite n ≤ {N_MAX_TABLA}.")
    cfg.validar()
    wb = Workbook()
    usa_mc = mc is not None and "media_muestral" in mc.columns and len(mc) > 0
    _hoja_leeme(wb, usa_mc)
    _hoja_binomial(wb, int(binom["n"]), float(binom["p0"]), float(binom["p1"]), int(binom["c"]), binom["cola"])
    _hoja_curva(wb)
    _hoja_z(wb, cfg, beta_obj)
    if usa_mc:
        _hoja_mc(wb, mc, cfg)
    if ejemplo_datos is None:
        ejemplo_datos = np.random.default_rng(4).normal(cfg.mu0, cfg.sigma, int(min(max(cfg.n, 5), 200)))
    _hoja_datos(wb, cfg, ejemplo_datos)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------------------------ lectura

def _num(v):
    try:
        if v is None or isinstance(v, bool):
            return None
        return float(v)
    except (TypeError, ValueError):
        return None


def _texto(v) -> str | None:
    return str(v).strip().lower() if v is not None else None


def leer(archivo) -> dict:
    """Lee un .xlsx subido. Devuelve lo que encuentre:

    - "datos": observaciones (hoja Tus_datos, columna A; si no existe, la primera columna numérica).
    - "params_datos": μ0, σ, α y cola de la hoja Tus_datos, si son válidos.
    - "config_z": ConfigZ desde la hoja Prueba_Z, si sus entradas son válidas.
    - "binomial": dict(n, p0, p1, c, cola) desde la hoja Binomial, si es válida.
    - "avisos": lista de mensajes para mostrar.
    """
    contenido = archivo.read() if hasattr(archivo, "read") else archivo
    if not contenido:
        raise ValueError("El archivo está vacío.")
    try:
        wb = load_workbook(io.BytesIO(contenido), data_only=False, read_only=False)
    except Exception as e:  # noqa: BLE001
        raise ValueError("No pude abrir el archivo como Excel (.xlsx).") from e
    out: dict = {"avisos": []}

    if "Prueba_Z" in wb.sheetnames:
        ws = wb["Prueba_Z"]
        vals = [_num(ws[f"B{i}"].value) for i in range(4, 9)]
        cola = _texto(ws["B9"].value)
        try:
            if None in vals:
                raise ValueError("hay entradas vacías o no numéricas en B4:B8")
            cfg = pz.ConfigZ(vals[0], vals[1], vals[2], int(round(vals[3])), vals[4], cola or "")
            cfg.validar()
            out["config_z"] = cfg
        except ValueError as e:
            out["avisos"].append(f"Hoja Prueba_Z: no la usé ({e}).")

    if "Binomial" in wb.sheetnames:
        ws = wb["Binomial"]
        vals = [_num(ws[f"B{i}"].value) for i in range(4, 8)]
        cola = _texto(ws["B8"].value)
        try:
            if None in vals:
                raise ValueError("hay entradas vacías o no numéricas en B4:B7")
            d = {"n": int(round(vals[0])), "p0": vals[1], "p1": vals[2], "c": int(round(vals[3])), "cola": cola or ""}
            bn.validar(d["n"], d["p0"], d["c"], d["cola"])
            bn.validar(d["n"], d["p1"], d["c"], d["cola"])
            if not 2 <= d["n"] <= N_MAX_TABLA:
                raise ValueError(f"n debe estar entre 2 y {N_MAX_TABLA}")
            if not bn.alternativa_valida(d["p0"], d["p1"], d["cola"]):
                raise ValueError("p1 no pertenece a la hipótesis alternativa")
            out["binomial"] = d
        except ValueError as e:
            out["avisos"].append(f"Hoja Binomial: no la usé ({e}).")

    datos = None
    if "Tus_datos" in wb.sheetnames:
        ws = wb["Tus_datos"]
        col = [_num(c.value) for (c,) in ws.iter_rows(min_row=2, min_col=1, max_col=1)]
        datos = np.array([v for v in col if v is not None])
        p = [_num(ws[f"D{i}"].value) for i in range(4, 7)]
        cola = _texto(ws["D7"].value)
        if None not in p and cola in pz.COLAS and p[1] > 0 and 0 < p[2] < 1:
            out["params_datos"] = {"mu0": p[0], "sigma": p[1], "alfa": p[2], "cola": cola}
    else:
        hoja = wb.worksheets[0]
        filas = list(hoja.iter_rows(values_only=True))
        if filas:
            _tabla(pd.DataFrame(filas[1:], columns=_nombres(filas[0])), out, f"la hoja «{hoja.title}»")
            datos = out.get("datos")
    if datos is not None and len(datos) >= 2:
        out["datos"] = datos
    elif datos is not None:
        out["avisos"].append("Encontré menos de 2 observaciones numéricas; no hay datos para la prueba.")
    return out


def _nombres(cabecera) -> list[str]:
    nombres, vistos = [], set()
    for j, v in enumerate(cabecera):
        base = str(v).strip() if v not in (None, "") else f"columna {j + 1}"
        nombre, i = base, 2
        while nombre in vistos:
            nombre, i = f"{base} ({i})", i + 1
        vistos.add(nombre)
        nombres.append(nombre)
    return nombres


def _tabla(df: pd.DataFrame, out: dict, origen: str) -> None:
    """Guarda la tabla y sus columnas numéricas; toma la primera como datos por defecto."""
    numericas = []
    for col in df.columns:
        if len(columna(df, col)) >= 2:
            numericas.append(col)
    out["tabla"] = df
    out["columnas"] = numericas
    if numericas:
        out["datos"] = columna(df, numericas[0])
        if len(numericas) > 1:
            out["avisos"].append(f"Encontré {len(numericas)} columnas numéricas en {origen}; elige cuál usar.")
    else:
        out["avisos"].append(f"No encontré columnas con al menos 2 números en {origen}.")


def columna(df: pd.DataFrame, nombre: str) -> np.ndarray:
    serie = df[nombre]
    if serie.dtype == object or str(serie.dtype).startswith(("string", "str")):  # números con coma decimal, como en descargas del DANE o BanRep
        serie = serie.astype(str).str.replace(".", "", regex=False).str.replace(",", ".", regex=False) \
            if serie.astype(str).str.contains(",").any() else serie
    return pd.to_numeric(serie, errors="coerce").dropna().to_numpy(dtype=float)


def leer_csv(contenido: bytes) -> dict:
    """CSV con separador automático (coma o punto y coma)."""
    out: dict = {"avisos": []}
    texto = None
    for cod in ("utf-8-sig", "latin-1"):
        try:
            texto = contenido.decode(cod)
            break
        except UnicodeDecodeError:
            continue
    if not texto or not texto.strip():
        raise ValueError("El archivo CSV está vacío.")
    df = pd.read_csv(io.StringIO(texto), sep=None, engine="python")
    _tabla(df, out, "el CSV")
    return out


def plantilla(cfg: pz.ConfigZ) -> bytes:
    """Libro mínimo con solo la hoja Tus_datos."""
    wb = Workbook()
    wb.remove(wb.active)
    ejemplo = np.random.default_rng(4).normal(cfg.mu0, cfg.sigma, int(min(max(cfg.n, 5), 200)))
    _hoja_datos(wb, cfg, ejemplo)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
